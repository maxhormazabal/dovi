from __future__ import annotations

import importlib.util
from pathlib import Path

from dovi.backend.abstract_backend import AbstractRenderBackend
from dovi.datamodel.base_models import OutputFormat
from dovi.datamodel.document import DoviDocument, TableItem, TextItem
from dovi.datamodel.results import RenderedArtifact
from dovi.exceptions import BackendNotAvailableError


class DocxRenderBackend(AbstractRenderBackend):
    output_format = OutputFormat.DOCX
    media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    extension = "docx"

    @classmethod
    def is_available(cls) -> bool:
        return importlib.util.find_spec("docx") is not None

    def render(self, document: DoviDocument, output_path: Path) -> RenderedArtifact:
        if not self.is_available():
            raise BackendNotAvailableError("docx", extra="docx")

        import docx

        out = docx.Document()
        for page_idx, page in enumerate(document.pages):
            if page_idx > 0:
                out.add_page_break()
            for item in document.iterate_items(page.page_no):
                if isinstance(item, TextItem):
                    if item.label == "title":
                        out.add_heading(item.text, level=0)
                    elif item.label == "section_header":
                        out.add_heading(item.text, level=min(item.level, 9))
                    else:
                        out.add_paragraph(item.text)
                elif isinstance(item, TableItem):
                    n_cols = max([len(item.header), *(len(r) for r in item.rows)], default=0)
                    if n_cols == 0:
                        continue
                    rows = ([item.header] if item.header else []) + item.rows
                    table = out.add_table(rows=len(rows), cols=n_cols)
                    table.style = "Table Grid"
                    for r, row in enumerate(rows):
                        for c, value in enumerate(row):
                            table.cell(r, c).text = value
        out.save(str(output_path))
        return self._artifact(output_path)
